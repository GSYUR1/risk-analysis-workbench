"""标准分析结果Excel导出库。

取数和ODPS认证不在本模块实现；分析数据由odps-query和ph_data_analysis提供。
"""

import re
import warnings

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")
pd.options.display.max_rows = 1000
pd.options.display.max_columns = 500
# 导出带格式的透视表到 Excel
def merge_formatted_pivot_to_excel(dict_df, output_path, risk_ind=None, sheet_name='all'):
    """
    将多个pivot_table写入同一个sheet，支持表头合并、数字格式化和条件格式
    
    参数: 
    dict_df: 字典，格式为{key: df}，key为标题，df为DataFrame（pivot_table，columns为MultiIndex）
    output_path: 输出路径
    risk_ind: 风险指标名称，如'3_30'，用于识别风险相关字段
    sheet_name: 唯一sheet名称，所有表纵向排列在该sheet中
    """ 
    
    wb = Workbook() 
    ws = wb.active
    ws.title = sheet_name
    
    # 样式定义
    header_align = Alignment(horizontal='center', vertical='center') 
    bold_font = Font(bold=True) 
    header_fill = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')
    bottom_border = Border(bottom=Side(style='thin')) 
    left_border = Border(left=Side(style='thin')) 
    left_bottom_border = Border(left=Side(style='thin'), bottom=Side(style='thin')) 
    
    def get_number_format(indicator): 
        """根据指标名返回对应的Excel数字格式""" 
        if indicator in ['计数', '平均授信', '平均动支', '成交人数', '件均', '成交件数', '成交金额']: 
            return '#,##0' 
        if indicator in ['占比', '金额占比', 'H24动支率', 'D7动支率', '动支率']: 
            return '0.00%' 
        if indicator in ['放款件均($)', '平均授信额度($)', '授信件均($)', '授信效率($)']: 
            return '#,##0.00' 
        if risk_ind: 
            if indicator in (f'm{risk_ind}', f'm{risk_ind}cnt', f'dpd{risk_ind}', f'dpd{risk_ind}cnt'): 
                return '0.00%' 
            if indicator in (f'm{risk_ind}lift', f'dpd{risk_ind}lift'): 
                return '0.00' 
            if indicator in (f'{risk_ind}出账', f'出账{risk_ind}'): 
                return '#,##0' 
        if re.match(r'^((m|dpd)\d+_\d+(cnt)?|cycle\d+_ever\d+)$', indicator): 
            return '0.00%' 
        if re.match(r'^((m|dpd)\d+_\d+|cycle\d+_ever\d+)lift$', indicator): 
            return '0.00' 
        if re.match(r'^(\d+_\d+出账|出账\d+_\d+|cycle\d+_ever\d+出账)$', indicator): 
            return '#,##0' 
        if indicator in ('IRR', 'APR') or re.search(r'率|占比', str(indicator)): 
            return '0.00%' 
        if str(indicator).endswith('($)') or indicator in ('久期', '名义期限'): 
            return '#,##0.00' 
        return None
    
    def get_cond_format_type(indicator): 
        """根据指标名返回条件格式类型""" 
        if indicator in ['占比', '金额占比', '成交人数', '计数']: 
            return 'databar_blue' 
        if indicator in ['平均授信', '平均动支', 'H24动支率', 'D7动支率', '动支率', '件均',
                         '放款件均($)', '平均授信额度($)']: 
            return 'colorscale_rwb' 
        if risk_ind: 
            if indicator in (f'm{risk_ind}', f'm{risk_ind}cnt', f'dpd{risk_ind}', f'dpd{risk_ind}cnt'): 
                return 'colorscale_ryg' 
            if indicator in (f'm{risk_ind}lift', f'dpd{risk_ind}lift'): 
                return 'colorscale_rw' 
            if indicator in (f'{risk_ind}出账', f'出账{risk_ind}'): 
                return 'highlight_green' 
        if re.match(r'^((m|dpd)\d+_\d+(cnt)?|cycle\d+_ever\d+)$', indicator): 
            return 'colorscale_ryg' 
        if re.match(r'^((m|dpd)\d+_\d+|cycle\d+_ever\d+)lift$', indicator): 
            return 'colorscale_rw' 
        if re.match(r'^(\d+_\d+出账|出账\d+_\d+|cycle\d+_ever\d+出账)$', indicator): 
            return 'highlight_green' 
        if re.search(r'通过率|转化率|动支率', str(indicator)): 
            return 'colorscale_rwb' 
        return None
    
    def build_range_excluding_cols(col_start, col_end, exclude_cols, row_start, row_end): 
        """构建排除指定列的范围字符串，用于条件格式整体应用""" 
        range_parts = [] 
        segment_start = None
        
        for col_num in range(col_start, col_end + 1): 
            if col_num not in exclude_cols: 
                if segment_start is None: 
                    segment_start = col_num
            else: 
                if segment_start is not None: 
                    range_parts.append((segment_start, col_num - 1)) 
                    segment_start = None
        
        if segment_start is not None: 
            range_parts.append((segment_start, col_end)) 
        
        if not range_parts: 
            return None
        
        return ' '.join([ 
            f"{get_column_letter(s)}{row_start}:{get_column_letter(e)}{row_end}" 
            for s, e in range_parts
        ]) 
    
    start_row = 1
    start_col = 1
    
    for key, df in dict_df.items():
        # 写入标题行
        title_cell = ws.cell(row=start_row, column=start_col, value=key)
        title_cell.font = Font(bold=True, size=11)
        title_cell.alignment = Alignment(horizontal='left', vertical='center')
        
        # 表头从标题行下一行开始
        start_row += 1
        
        # Styler 或带展示包装的对象：尽量取回底层 DataFrame
        if not isinstance(df, pd.DataFrame):
            df = getattr(df, 'data', df)
        df = pd.DataFrame(df)

        # 列索引层数不限：level0=指标名，其余为列维度（col 传 list 时会有多层）
        col_tuples = [c if isinstance(c, tuple) else (c,) for c in df.columns]
        n_col_levels = max(len(t) for t in col_tuples)
        col_tuples = [t + ('',) * (n_col_levels - len(t)) for t in col_tuples]
        level0 = [t[0] for t in col_tuples]

        # 行索引层数不限：ind 传 list 时每层占一列
        n_idx = df.index.nlevels
        idx_names = [n if n is not None else '' for n in df.index.names]

        first_data_col = start_col + n_idx
        header_row1 = start_row
        header_row_last = header_row1 + n_col_levels - 1
        data_start_row = header_row_last + 1

        # === 1. 各层表头的合并区间（按 0..level 前缀切段，同名不同父不会误合并）===
        level_segments = []
        for level in range(n_col_levels):
            segments = []
            prev_key, seg_start = None, 0
            for col_idx, t in enumerate(col_tuples):
                key = t[:level + 1]
                if prev_key is None:
                    prev_key, seg_start = key, col_idx
                elif key != prev_key:
                    segments.append((prev_key[level], seg_start, col_idx - 1))
                    prev_key, seg_start = key, col_idx
            if prev_key is not None:
                segments.append((prev_key[level], seg_start, len(col_tuples) - 1))
            level_segments.append(segments)

        # 指标块（level0）区间，用于左边框和条件格式
        merge_info = [(label, first_data_col + s, first_data_col + e)
                      for label, s, e in level_segments[0]]
        block_start_cols = {c_start for _, c_start, _ in merge_info}

        # All 汇总列：列维度任一层为 All
        all_col_indices = {
            first_data_col + col_idx
            for col_idx, t in enumerate(col_tuples)
            if any(str(x) == 'All' for x in t[1:])
        }

        # === 2. 逐层写入列表头并合并 ===
        for level, segments in enumerate(level_segments):
            row_num = header_row1 + level
            is_last_level = (level == n_col_levels - 1)
            for label, s, e in segments:
                c_start = first_data_col + s
                c_end = first_data_col + e
                cell = ws.cell(row=row_num, column=c_start, value=_indicator_excel_value(label))
                cell.alignment = header_align
                cell.font = bold_font
                if c_end > c_start:
                    ws.merge_cells(start_row=row_num, start_column=c_start,
                                   end_row=row_num, end_column=c_end)
                for col in range(c_start, c_end + 1):
                    hc = ws.cell(row=row_num, column=col)
                    hc.fill = header_fill
                    if col in block_start_cols:
                        hc.border = left_bottom_border if is_last_level else left_border
                    elif is_last_level:
                        hc.border = bottom_border

        # === 3. 写入索引列表头（每层一列，纵向合并到表头最后一行）===
        for i, name in enumerate(idx_names):
            col_num = start_col + i
            cell = ws.cell(row=header_row1, column=col_num, value=name)
            cell.font = bold_font
            cell.alignment = header_align
            if n_col_levels > 1:
                ws.merge_cells(start_row=header_row1, start_column=col_num,
                               end_row=header_row_last, end_column=col_num)
            for r in range(header_row1, header_row_last + 1):
                ws.cell(row=r, column=col_num).fill = header_fill
            ws.cell(row=header_row_last, column=col_num).border = bottom_border

        # === 4. 写入数据并应用格式 ===
        all_row_offsets = []
        for row_idx, (idx, row_data) in enumerate(df.iterrows()):
            row_num = data_start_row + row_idx
            is_all_row = _indicator_is_all_row_label(idx)
            if is_all_row:
                all_row_offsets.append(row_idx)

            # 索引列（多层索引每层写一列）
            idx_vals = list(idx) if isinstance(idx, tuple) else [idx]
            while len(idx_vals) < n_idx:
                idx_vals.append('')
            for i, iv in enumerate(idx_vals):
                idx_cell = ws.cell(row=row_num, column=start_col + i,
                                   value=_indicator_excel_value(iv))
                if is_all_row:
                    idx_cell.font = bold_font

            # 数据列
            for col_idx, val in enumerate(row_data):
                col_num = first_data_col + col_idx
                cell = ws.cell(row=row_num, column=col_num, value=_indicator_excel_value(val))

                if col_num in block_start_cols:
                    cell.border = left_border

                fmt = get_number_format(level0[col_idx])
                if fmt:
                    cell.number_format = fmt

                if is_all_row or col_num in all_col_indices:
                    cell.font = bold_font

        # === 5. 应用条件格式（整体应用，排除All行和All列）===
        n_rows = len(df)
        non_all_rows = [i for i in range(n_rows) if i not in all_row_offsets]
        row_segments = []
        for i in non_all_rows:
            if row_segments and i == row_segments[-1][1] + 1:
                row_segments[-1][1] = i
            else:
                row_segments.append([i, i])

        for indicator, c_start, c_end in merge_info:
            cond_type = get_cond_format_type(indicator)
            if not cond_type:
                continue

            for r0, r1 in row_segments:
                # 构建排除All列的范围字符串
                range_str = build_range_excluding_cols(
                    c_start, c_end, all_col_indices,
                    data_start_row + r0, data_start_row + r1
                )
                if not range_str:
                    continue

                if cond_type == 'databar_blue':
                    rule = DataBarRule(start_type='min', end_type='max',
                                       color='5B9BD5', showValue=True)
                elif cond_type == 'colorscale_rwb':
                    rule = ColorScaleRule(start_type='min', start_color='5A8AC6',
                                          mid_type='percentile', mid_value=50, mid_color='FFFFFF',
                                          end_type='max', end_color='F8696B')
                elif cond_type == 'colorscale_ryg':
                    rule = ColorScaleRule(start_type='min', start_color='63BE7B',
                                          mid_type='percentile', mid_value=50, mid_color='FFEB84',
                                          end_type='max', end_color='F8696B')
                elif cond_type == 'colorscale_rw':
                    rule = ColorScaleRule(start_type='min', start_color='FFFFFF',
                                          end_type='max', end_color='F8696B')
                elif cond_type == 'highlight_green':
                    green_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
                    rule = CellIsRule(operator='greaterThan', formula=['500'], fill=green_fill)

                ws.conditional_formatting.add(range_str, rule)

        # 更新下一个表格的起始行（标题行 + 表头 + 数据 + 间隔）
        start_row = data_start_row + n_rows + 3

    wb.save(output_path)


# 兼容旧函数名
merge_formatted_excel_func = merge_formatted_pivot_to_excel


# generate_indicator 数据条颜色（与 ph_data_analysis._style_indicator_df 对齐）
_INDICATOR_BAR_COLOR_COUNT = 'AED0EA'       # 雾蓝 — 申请/件数
_INDICATOR_BAR_COLOR_PASS_RATE = 'B5D8A8'   # 抹茶 — 通过率/转化率
_INDICATOR_BAR_COLOR_DRAW_RATE = 'F3C89A'   # 浅杏 — 动支率/定价
_INDICATOR_BAR_COLOR_AMOUNT = 'E8D89A'      # 香槟金 — 金额/件均
_INDICATOR_BAR_COLOR_RISK = 'EBB0B0'        # 浅珊瑚 — 风险

_RE_INDICATOR_RISK_COL = re.compile(r'^(dpd\d+_\d+(lift|cnt)?|出账\d+_\d+|\d+_\d+出账|m\d+_\d+(lift|cnt)?)$')
_RE_INDICATOR_DPD_RATE = re.compile(r'^(dpd|m)\d+_\d+(cnt)?$')
_RE_INDICATOR_DPD_LIFT = re.compile(r'^(dpd|m)\d+_\d+lift$')


def _indicator_is_all_row_label(label):
    """判断是否为汇总行：indicator 的 ALL。"""
    if isinstance(label, tuple):
        tokens = [str(x) for x in label]
        return any(t in ('ALL', 'All') for t in tokens) and all(t in ('ALL', 'All', '') for t in tokens)
    return str(label) in ('ALL', 'All')


def _indicator_excel_value(val):
    """把 pandas/numpy 值转成 openpyxl 可写的标量。"""
    if val is None or (isinstance(val, float) and np.isnan(val)):
        return None
    try:
        if pd.isna(val):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(val, (np.integer,)):
        return int(val)
    if isinstance(val, (np.floating,)):
        return float(val)
    if isinstance(val, (np.bool_,)):
        return bool(val)
    return val


def _indicator_number_format(name):
    """与 ph_data_analysis._metric_number_format 对齐的 Excel 数字格式。"""
    name = str(name)
    if _RE_INDICATOR_DPD_LIFT.match(name):
        return '0.00'
    if name in ('IRR', 'APR') or re.search(r'率|占比', name) or _RE_INDICATOR_DPD_RATE.match(name):
        return '0.00%'
    if name.endswith('($)') or name in ('久期', '名义期限'):
        return '#,##0.00'
    return '#,##0'


def _indicator_bar_color(name):
    """按指标语义返回数据条颜色（无 #），与 generate_indicator Jupyter 着色一致。"""
    name = str(name)
    if _RE_INDICATOR_RISK_COL.match(name):
        return _INDICATOR_BAR_COLOR_RISK
    if '动支率' in name or name == '动支pending率' or name in ('久期', '名义期限', 'IRR', 'APR'):
        return _INDICATOR_BAR_COLOR_DRAW_RATE
    if re.search(r'通过率|转化率', name):
        return _INDICATOR_BAR_COLOR_PASS_RATE
    if re.search(r'件均|效率|金额|额度', name):
        return _INDICATOR_BAR_COLOR_AMOUNT
    return _INDICATOR_BAR_COLOR_COUNT


def merge_formatted_indicator_to_excel(dict_df, output_path, sheet_name='all'):
    """
    将多个 generate_indicator 结果表写入同一个 sheet，带数字格式和按列类型的数据条。

    对齐 Jupyter 中 generate_indicator(show=True) 的展示：
    - 百分比 / 千分位 / 金额小数格式按列名规则
    - 件数/通过率/动支率/金额/风险分色数据条（vmin=0，按列各自缩放）
    - ALL 汇总行加粗、不画数据条

    参数:
    dict_df: 字典 {标题: DataFrame}，value 为 generate_indicator 返回的表（单层列，索引可为单层或多层）
    output_path: Excel 保存路径
    sheet_name: 工作表名称，所有表纵向排列
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_align = Alignment(horizontal='center', vertical='center')
    bold_font = Font(bold=True)
    header_fill = PatternFill(start_color='DCE6F1', end_color='DCE6F1', fill_type='solid')
    bottom_border = Border(bottom=Side(style='thin'))

    start_row = 1
    start_col = 1

    for key, df in dict_df.items():
        if df is None or len(df.columns) == 0:
            continue

        # Styler 或带展示包装的对象：尽量取回底层 DataFrame
        if not isinstance(df, pd.DataFrame):
            df = getattr(df, 'data', df)
        df = pd.DataFrame(df)

        title_cell = ws.cell(row=start_row, column=start_col, value=key)
        title_cell.font = Font(bold=True, size=11)
        title_cell.alignment = Alignment(horizontal='left', vertical='center')
        start_row += 1

        if isinstance(df.columns, pd.MultiIndex):
            col_names = ['_'.join([str(x) for x in c if str(x) != '']) for c in df.columns]
        else:
            col_names = list(df.columns)

        if isinstance(df.index, pd.MultiIndex):
            n_idx = df.index.nlevels
            idx_names = [n if n is not None else '' for n in df.index.names]
        else:
            n_idx = 1
            idx_names = [df.index.name if df.index.name else '']

        header_row = start_row
        data_start_row = header_row + 1

        for i, name in enumerate(idx_names):
            cell = ws.cell(row=header_row, column=start_col + i, value=name)
            cell.font = bold_font
            cell.alignment = header_align
            cell.fill = header_fill
            cell.border = bottom_border

        for col_idx, name in enumerate(col_names):
            cell = ws.cell(row=header_row, column=start_col + n_idx + col_idx, value=name)
            cell.font = bold_font
            cell.alignment = header_align
            cell.fill = header_fill
            cell.border = bottom_border

        all_row_offsets = []
        for row_idx, (idx, row_data) in enumerate(df.iterrows()):
            row_num = data_start_row + row_idx
            is_all_row = _indicator_is_all_row_label(idx)
            if is_all_row:
                all_row_offsets.append(row_idx)

            if isinstance(idx, tuple):
                idx_vals = list(idx)
            else:
                idx_vals = [idx]
            while len(idx_vals) < n_idx:
                idx_vals.append('')

            for i, iv in enumerate(idx_vals):
                idx_cell = ws.cell(row=row_num, column=start_col + i, value=_indicator_excel_value(iv))
                if is_all_row:
                    idx_cell.font = bold_font

            for col_idx, val in enumerate(row_data):
                col_num = start_col + n_idx + col_idx
                cell = ws.cell(row=row_num, column=col_num, value=_indicator_excel_value(val))
                fmt = _indicator_number_format(col_names[col_idx])
                if fmt:
                    cell.number_format = fmt
                if is_all_row:
                    cell.font = bold_font

        # 数据条：排除 ALL 行；每列单独一条规则（与 generate_indicator 按列缩放一致）
        n_rows = len(df)
        if n_rows > 0:
            non_all = [i for i in range(n_rows) if i not in all_row_offsets]
            if non_all:
                # ALL 夹在中间时拆成连续段
                segments = []
                seg_s = non_all[0]
                prev = non_all[0]
                for i in non_all[1:]:
                    if i == prev + 1:
                        prev = i
                    else:
                        segments.append((data_start_row + seg_s, data_start_row + prev))
                        seg_s = i
                        prev = i
                segments.append((data_start_row + seg_s, data_start_row + prev))

                for col_idx, name in enumerate(col_names):
                    series = df.iloc[:, col_idx]
                    if not series.notna().any():
                        continue
                    color = _indicator_bar_color(name)
                    col_letter = get_column_letter(start_col + n_idx + col_idx)
                    for r0, r1 in segments:
                        if r1 < r0:
                            continue
                        range_str = f'{col_letter}{r0}:{col_letter}{r1}'
                        rule = DataBarRule(
                            start_type='num', start_value=0,
                            end_type='max',
                            color=color, showValue=True
                        )
                        ws.conditional_formatting.add(range_str, rule)

        start_row = data_start_row + n_rows + 3

    wb.save(output_path)



